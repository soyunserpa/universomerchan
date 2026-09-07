"""QA renderer: approximates a PPTX with HTML/CSS and screenshots each slide (fallback when soffice is unavailable)."""
import sys, base64, io
from pptx import Presentation
from pptx.util import Emu
from pptx.enum.shapes import MSO_SHAPE_TYPE
from playwright.sync_api import sync_playwright

SRC = sys.argv[1]; OUT = sys.argv[2]
PX = 100  # px per inch
prs = Presentation(SRC)
SW, SH = Emu(prs.slide_width).inches*PX, Emu(prs.slide_height).inches*PX

def rgb(c):
    try:
        return '#'+str(c.rgb)
    except Exception:
        return None

def shape_html(sh):
    x, y = Emu(sh.left).inches*PX, Emu(sh.top).inches*PX
    w, h = Emu(sh.width).inches*PX, Emu(sh.height).inches*PX
    out = []
    if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
        b = base64.b64encode(sh.image.blob).decode()
        out.append(f'<img src="data:{sh.image.content_type};base64,{b}" style="position:absolute;left:{x}px;top:{y}px;width:{w}px;height:{h}px;object-fit:fill">')
        return ''.join(out)
    style = f'position:absolute;left:{x}px;top:{y}px;width:{w}px;height:{h}px;box-sizing:border-box;'
    geom = ''
    try:
        prst = sh._element.spPr.prstGeom.get('prst')
    except Exception:
        prst = None
    if prst == 'roundRect':
        # radius = adj * min(w,h); pptxgenjs writes adj in 1/100000 of shape; approximate
        adj = None
        try:
            gd = sh._element.spPr.prstGeom.avLst.findall('{http://schemas.openxmlformats.org/drawingml/2006/main}gd')
            if gd: adj = int(gd[0].get('fmla').split()[-1])/100000.0
        except Exception: pass
        r = (adj if adj is not None else 0.1667)*min(w,h)
        geom = f'border-radius:{r}px;'
    elif prst == 'ellipse':
        geom = 'border-radius:50%;'
    fill = None
    try:
        if sh.fill.type == 1: fill = rgb(sh.fill.fore_color)
    except Exception: pass
    line = None
    try:
        if sh.line.fill.type == 1:
            line = rgb(sh.line.color); lw = sh.line.width.pt if sh.line.width else 0.75
    except Exception: pass
    if prst == 'line':
        col = line or '#000'
        out.append(f'<div style="{style}border-top:{max(1,lw)}px solid {col};height:0"></div>')
        return ''.join(out)
    if fill: style += f'background:{fill};'
    if line: style += f'border:{lw}px solid {line};'
    if fill or line:
        out.append(f'<div style="{style}{geom}"></div>')
    if sh.has_text_frame and sh.text_frame.text.strip():
        tf = sh.text_frame
        bp = tf._txBody.bodyPr
        anchor = bp.get('anchor') or 't'
        va = {'t':'flex-start','ctr':'center','b':'flex-end'}.get(anchor,'flex-start')
        ins = lambda a,d: (int(bp.get(a))/12700 if bp.get(a) is not None else d)
        pl, pr, pt, pb = ins('lIns',7.2), ins('rIns',7.2), ins('tIns',3.6), ins('bIns',3.6)
        paras = []
        for para in tf.paragraphs:
            al = {1:'left',2:'center',3:'right'}.get(para.alignment, 'left') if para.alignment else 'left'
            ls = ''
            try:
                if para.line_spacing is not None and hasattr(para.line_spacing,'pt'):
                    ls = f'line-height:{para.line_spacing.pt*PX/72}px;'
            except Exception: pass
            bullet = ''
            pPr = para._p.pPr
            if pPr is not None and pPr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}buChar') is not None:
                bullet = '&bull; '
            runs = []
            for r in para.runs:
                f = r.font
                fs = (f.size.pt if f.size else 18)*PX/72
                col = rgb(f.color) if f.color and f.color.type is not None else '#000'
                fam = f.name or 'Poppins'
                weight = 700 if f.bold else (300 if 'Light' in fam else 400)
                fam = "'Poppins'"
                st = f'font-family:{fam},sans-serif;font-size:{fs}px;font-weight:{weight};color:{col};'
                if f.italic: st += 'font-style:italic;'
                txt = (r.text or '').replace('&','&amp;').replace('<','&lt;')
                runs.append(f'<span style="{st}">{txt}</span>')
            if not runs and para.text == '':
                runs.append('<span style="font-size:10px">&nbsp;</span>')
            paras.append(f'<div style="text-align:{al};{ls}white-space:pre-wrap">{bullet}{"".join(runs)}</div>')
        out.append(f'<div style="{style}display:flex;flex-direction:column;justify-content:{va};padding:{pt}px {pr}px {pb}px {pl}px;overflow:visible">{"".join(paras)}</div>')
    return ''.join(out)

pages = []
for si, slide in enumerate(prs.slides):
    bg = '#ffffff'
    try:
        if slide.background.fill.type == 1: bg = rgb(slide.background.fill.fore_color)
    except Exception: pass
    body = ''.join(shape_html(sh) for sh in slide.shapes)
    html = f"""<html><head><link href="https://fonts.googleapis.com/css2?family=Poppins:ital,wght@0,300;0,400;0,700;1,300;1,400&display=swap" rel="stylesheet">
    <style>body{{margin:0;width:{SW}px;height:{SH}px;overflow:hidden;background:{bg};position:relative}}</style></head><body>{body}</body></html>"""
    pages.append(html)

import glob, os
exe = glob.glob('/opt/pw-browsers/chromium_headless_shell-*/chrome-linux/headless_shell')[0]
with sync_playwright() as p:
    b = p.chromium.launch(executable_path=exe); pg = b.new_page(viewport={'width':int(SW),'height':int(SH)})
    for i, html in enumerate(pages):
        pg.set_content(html); pg.wait_for_timeout(700)
        pg.screenshot(path=f'{OUT}/slide{i+1}.png'); print('slide', i+1)
    b.close()
